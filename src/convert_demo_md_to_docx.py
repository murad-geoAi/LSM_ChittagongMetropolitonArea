"""
Convert the demo manuscript from Markdown to DOCX.

This is a lightweight Markdown-to-Word converter tailored to the structure used in
`docs/DEMO_RESEARCH_PAPER.md` (headings, paragraphs, lists, tables, and images).

Run (from LSM_ChittagongMetropolitonArea):
    python src/convert_demo_md_to_docx.py

Optional:
    python src/convert_demo_md_to_docx.py --input docs/DEMO_RESEARCH_PAPER.md --output docs/DEMO_RESEARCH_PAPER_FROM_MD.docx
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from urllib.parse import unquote

from docx import Document
from docx.shared import Inches, Pt

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
IMAGE_RE = re.compile(r"!\[(?P<alt>.*?)\]\((?P<path>.*?)\)")
BULLET_RE = re.compile(r"^\s*-\s+(.*)$")
NUMBERED_RE = re.compile(r"^\s*\d+\.\s+(.*)$")
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?\s*$")


def clean_inline_markdown(text: str) -> str:
    """
    Remove/normalize a small subset of inline Markdown to plain text.

    This keeps the conversion dependency-free (no pandoc/markdown renderer) and is
    adequate for this repo’s manuscript formatting.
    """

    text = text.replace("\\_", "_")
    text = text.replace("\\*", "*")
    text = text.replace("\\`", "`")
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)  # links -> label
    text = text.replace("**", "")
    text = text.replace("*", "")
    text = text.replace("`", "")
    return text.strip()


def is_table_header(line: str, next_line: str | None) -> bool:
    if next_line is None:
        return False
    if "|" not in line:
        return False
    return bool(TABLE_SEP_RE.match(next_line))


def parse_table_row(line: str) -> list[str]:
    row = line.strip().strip("|")
    cells = [clean_inline_markdown(c.strip()) for c in row.split("|")]
    return cells


def add_table(document: Document, header_line: str, row_lines: list[str]) -> None:
    header_cells = parse_table_row(header_line)
    rows = [parse_table_row(r) for r in row_lines]
    col_count = max([len(header_cells)] + [len(r) for r in rows] + [1])

    table = document.add_table(rows=len(rows) + 1, cols=col_count)
    table.style = "Table Grid"

    for j, cell_text in enumerate(header_cells):
        table.cell(0, j).text = cell_text

    for i, row in enumerate(rows, start=1):
        for j, cell_text in enumerate(row):
            table.cell(i, j).text = cell_text


def add_code_block(document: Document, code_lines: list[str]) -> None:
    if not code_lines:
        return
    paragraph = document.add_paragraph()
    run = paragraph.add_run("\n".join(code_lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def add_image(document: Document, md_dir: Path, alt_text: str, rel_path: str) -> None:
    caption = clean_inline_markdown(alt_text)
    img_rel = Path(unquote(rel_path))
    img_path = (md_dir / img_rel).resolve()

    if img_path.exists():
        document.add_picture(str(img_path), width=Inches(6.2))
        cap_para = document.add_paragraph(caption)
        try:
            cap_para.style = "Caption"
        except Exception:
            pass
    else:
        document.add_paragraph(f"[Missing image: {rel_path}] {caption}")


def convert_markdown_to_docx(md_path: Path, docx_path: Path) -> None:
    md_dir = md_path.parent
    lines = md_path.read_text(encoding="utf-8").splitlines()

    doc = Document()

    in_code = False
    code_lines: list[str] = []
    first_heading_written = False

    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip("\n")

        if line.strip().startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        heading_match = HEADING_RE.match(line)
        if heading_match:
            level = len(heading_match.group(1))
            text = clean_inline_markdown(heading_match.group(2))
            if level == 1 and not first_heading_written:
                doc.add_heading(text, level=0)
                first_heading_written = True
            else:
                doc.add_heading(text, level=min(level, 9))
            i += 1
            continue

        image_match = IMAGE_RE.fullmatch(line.strip())
        if image_match:
            add_image(
                doc,
                md_dir=md_dir,
                alt_text=image_match.group("alt"),
                rel_path=image_match.group("path"),
            )
            i += 1
            continue

        next_line = lines[i + 1] if i + 1 < len(lines) else None
        if is_table_header(line, next_line):
            header_line = line
            i += 2  # skip header + separator
            row_lines: list[str] = []
            while i < len(lines) and lines[i].strip():
                if "|" not in lines[i]:
                    break
                row_lines.append(lines[i])
                i += 1
            add_table(doc, header_line, row_lines)
            continue

        bullet_match = BULLET_RE.match(line)
        if bullet_match:
            while i < len(lines):
                m = BULLET_RE.match(lines[i])
                if not m:
                    break
                doc.add_paragraph(clean_inline_markdown(m.group(1)), style="List Bullet")
                i += 1
            continue

        numbered_match = NUMBERED_RE.match(line)
        if numbered_match:
            while i < len(lines):
                m = NUMBERED_RE.match(lines[i])
                if not m:
                    break
                doc.add_paragraph(clean_inline_markdown(m.group(1)), style="List Number")
                i += 1
            continue

        # Paragraph: collect until blank line or next block marker.
        para_lines = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip():
            look = lines[i]
            if HEADING_RE.match(look):
                break
            if look.strip().startswith("```"):
                break
            if IMAGE_RE.fullmatch(look.strip()):
                break
            if BULLET_RE.match(look) or NUMBERED_RE.match(look):
                break
            if is_table_header(look, lines[i + 1] if i + 1 < len(lines) else None):
                break
            para_lines.append(look.strip())
            i += 1

        paragraph_text = clean_inline_markdown(" ".join(para_lines))
        doc.add_paragraph(paragraph_text)

    if in_code and code_lines:
        add_code_block(doc, code_lines)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--input",
        default="docs/DEMO_RESEARCH_PAPER.md",
        help="Input Markdown path (relative to repo root).",
    )
    parser.add_argument(
        "--output",
        default="docs/DEMO_RESEARCH_PAPER_FROM_MD.docx",
        help="Output DOCX path (relative to repo root).",
    )
    args = parser.parse_args()

    md_path = Path(args.input)
    docx_path = Path(args.output)

    if not md_path.exists():
        raise FileNotFoundError(f"Input Markdown not found: {md_path}")

    convert_markdown_to_docx(md_path=md_path, docx_path=docx_path)
    print(f"Generated {docx_path}")


if __name__ == "__main__":
    main()
