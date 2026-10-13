"""Meaningful release checks for numeric consistency, anonymization and DOCX integrity."""
from pathlib import Path
from zipfile import ZipFile
import json, re, hashlib, importlib.metadata
import numpy as np,pandas as pd
from docx import Document
from lxml import etree
import pymupdf
R=Path(__file__).resolve().parent;A=R/'analysis';S=R/'submission'
report={};checks=[]
def check(name,condition):
 if not bool(condition):raise AssertionError(name)
 checks.append(name)
def texts(d):return [p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells]
clean=Document(S/'Manuscript_Revised_Clean.docx');red=Document(S/'Manuscript_Revised_RedMarked.docx')
check('Clean and red copies have identical text and table content',texts(clean)==texts(red))
check('Five main figures and three tables',len(clean.inline_shapes)==5 and len(clean.tables)==3)
abstract=clean.paragraphs[2].text
check('Abstract length between 150 and 300 words',150<=len(abstract.split())<=300)
check('Ten point single spaced manuscript',clean.styles['Normal'].font.size.pt==10 and clean.styles['Normal'].paragraph_format.line_spacing==1)
for name in ['Manuscript_Revised_Clean.docx','Manuscript_Revised_RedMarked.docx','Supplementary_Methods_Results.docx','Response_to_Reviewers.docx']:
 d=Document(S/name);tx='\n'.join(texts(d));check(name+' no unfilled placeholders',not re.search(r'TODO|TBD|\[INSERT|AUTHOR COMPLETION REQUIRED',tx))
 if name.startswith('Manuscript_'):
  check(name+' anonymous main document',not any(x in tx for x in ['Golam Murad','Md. Aftabur','u1901087','murad-geoAi','yasuhara.hideaki']))
 with ZipFile(S/name) as z:
  x=etree.fromstring(z.read('word/document.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
  check(name+' no tracked changes',not x.findall('.//w:ins',ns) and not x.findall('.//w:del',ns))
  check(name+' no corrupted XML characters',b'\xef\xbf\xbd' not in z.read('word/document.xml'))
response=Document(S/'Response_to_Reviewers.docx');tx='\n'.join(texts(response))
check('All nine Reviewer 1 and eight Reviewer 2 responses',all(f'Reviewer 1 comment {i}' in tx for i in range(1,10)) and all(f'Reviewer 2 comment {i}' in tx for i in range(1,9)))
original=Document(R/'Partially correctde Manuscript_LSM_Murad_et_al_Revised_RedMarked.docx');title=Document(S/'Title_Page_AUTHOR_COMPLETION_REQUIRED.docx')
check('Original author block retained on title page',all(p.text in texts(title) for p in original.paragraphs[2:10] if p.text))
hold=pd.read_csv(A/'holdout_metrics_reproduced.csv');check('Confusion matrices total 74 for every model',((hold.TN+hold.FP+hold.FN+hold.TP)==74).all())
check('Every holdout has 31 positives',((hold.FN+hold.TP)==31).all())
splits=pd.read_csv(A/'spatial_assignments.csv');check('295 train and 74 holdout',splits.split.value_counts().to_dict()=={'training':295,'holdout':74})
check('Training samples each have one spatial test fold',splits[splits.split=='training'].spatial_test_fold.notna().all())
check('Holdout excluded from spatial CV',splits[splits.split=='holdout'].spatial_test_fold.isna().all())
unc=pd.read_csv(A/'corrected_gpc_point_uncertainty.csv')
check('Nonnegative variances', (unc[['Latent_variance','Probability_variance','Expected_Bernoulli_variance']]>=0).all().all())
check('Bernoulli variance decomposition',np.allclose(unc.Probability_variance+unc.Expected_Bernoulli_variance,unc.Quadrature_probability_mean*(1-unc.Quadrature_probability_mean),atol=1e-12))
check('Latent uncertainty distinct from old ambiguity',not np.allclose(unc.Normalized_latent_sd,1-2*abs(unc.Probability-.5)))
check('Tertiles contain 123 points each',unc.Susceptibility_tertile.value_counts().tolist()==[123,123,123] and unc.Uncertainty_tertile.value_counts().tolist()==[123,123,123])
source=R/'inputs/Landslide_Data.csv';audit=json.loads((A/'data_audit.json').read_text())
check('Input source hash unchanged',hashlib.sha256(source.read_bytes()).hexdigest()==audit['source_sha256'])
render=json.loads((R/'qa/render_report.json').read_text());pages={}
for item in render:
 pdf=R/'qa'/Path(item['document']).stem/(Path(item['document']).stem+'.pdf');doc=pymupdf.open(pdf);pages[item['document']]=len(doc)
 check(item['document']+' no blank PDF pages',all(len(p.get_text().strip())>30 or len(p.get_images())>0 for p in doc))
 check(item['document']+' rendered after latest authoring',pdf.stat().st_mtime>=(S/item['document']).stat().st_mtime)
 for p in doc:
  for b in p.get_text('blocks'):
   check(item['document']+f' page {p.number+1} text stays inside page',b[0]>=0 and b[1]>=0 and b[2]<=p.rect.width+.5 and b[3]<=p.rect.height+.5)
report.update(status='PASSED technical and document checks; scientific and author release conditions remain',checks=checks,abstract_words=len(abstract.split()),rendered_pages=pages,release_ready=False)
(R/'audit/package_verification.json').write_text(json.dumps(report,indent=2))
packages=['numpy','pandas','scikit-learn','scipy','matplotlib','shap','jenkspy','joblib','threadpoolctl','python-docx','lxml','Pillow','PyMuPDF','pywin32']
(R/'requirements-revision.txt').write_text('\n'.join(f'{p}=={importlib.metadata.version(p)}' for p in packages)+'\n')
print('PASSED',len(checks),'checks;',sum(pages.values()),'rendered pages; author/data release conditions remain.')
